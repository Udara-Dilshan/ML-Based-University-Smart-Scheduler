from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
from io import BytesIO
import tempfile
from pathlib import Path
import sys
import subprocess

from app.database.connection import get_db
from app.models.timetable import TimetableSession
from app.models.academic import Batch, BatchActiveTerm
from app.models.user import User
from app.utils.dependencies import get_current_user

try:
    import docx
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    docx = None

try:
    from docx2pdf import convert as docx_to_pdf
except ImportError:
    docx_to_pdf = None

router = APIRouter(prefix="/api/timetable/export", tags=["timetable-export"])

def _ensure_dependencies(format_name: str) -> None:
    if format_name == "docx" and docx is None:
        raise HTTPException(status_code=500, detail="python-docx is not installed")
    if format_name == "pdf":
        if docx is None:
            raise HTTPException(status_code=500, detail="python-docx is not installed")
        if sys.platform == "win32" and docx_to_pdf is None:
            raise HTTPException(status_code=500, detail="docx2pdf is not installed")

def _set_cell_shading(cell, color_hex: str):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color_hex)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def _add_centered_paragraph(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    p.paragraph_format.space_after = Pt(0)
    return p

def append_timetable_to_doc(doc, batch: Batch, active_term: Optional[BatchActiveTerm], sessions: list, is_first: bool = True):
    if not is_first:
        doc.add_page_break()

    # Headers
    degree_name = batch.degree.name if batch.degree else "Unknown Degree"
    dept_name = batch.degree.department.name if batch.degree and batch.degree.department else "Unknown Department"
    faculty_name = batch.degree.department.faculty.name if batch.degree and batch.degree.department and batch.degree.department.faculty else "Unknown Faculty"
    academic_year = batch.batch_code
    semester_name = active_term.semester_name if active_term else f"Semester {batch.current_semester}"
    
    level_str = ""
    try:
        # Assuming semester_name might be "Year 1 Semester 2" or we can guess from batch.current_semester
        year = ((batch.current_semester - 1) // 2) + 1
        level_str = f" - {year * 100} Level"
    except Exception:
        pass

    _add_centered_paragraph(doc, f"Time Table - {degree_name} Degree Programme", size=11)
    _add_centered_paragraph(doc, f"Department of {dept_name}", size=11)
    _add_centered_paragraph(doc, f"Faculty of {faculty_name}", size=11)
    _add_centered_paragraph(doc, f"Academic Year: {academic_year} - {semester_name}{level_str}", size=11)
    doc.add_paragraph()

    # Table 1: Timetable Grid
    # 13 rows: header + 12 hourly slots (8am-8pm)
    table = doc.add_table(rows=13, cols=6)
    table.style = 'Table Grid'
    table.autofit = False

    widths = [Inches(1.5), Inches(1.2), Inches(1.2), Inches(1.2), Inches(1.2), Inches(1.2)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    days = ['', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    for idx, day in enumerate(days):
        cell = table.cell(0, idx)
        cell.text = day
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.bold = True

    times = [
        '8.00 am - 9.00 am', '9.00 am - 10.00 am', '10.00 am - 11.00 am',
        '11.00 am - 12.00 noon', '12.00 noon - 1.00 pm', '1.00 pm - 2.00 pm',
        '2.00 pm - 3.00 pm', '3.00 pm - 4.00 pm', '4.00 pm - 5.00 pm',
        '5.00 pm - 6.00 pm', '6.00 pm - 7.00 pm', '7.00 pm - 8.00 pm'
    ]

    for idx, time_str in enumerate(times, start=1):
        cell = table.cell(idx, 0)
        cell.text = time_str
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9)

    # Identify lunch slot based on first session or default to 12.00 noon
    lunch_row = 5 # 12.00 noon - 1.00 pm
    lunch_start = 12 * 60
    if sessions:
        # Check if batch has a specific lunch break
        from app.models.settings import SystemConstraint
        db_session = Session.object_session(sessions[0]) if sessions else None
        if db_session:
            gc_start = db_session.query(SystemConstraint).filter_by(batch_id=batch.batch_id, name="lunch_break_start").first()
            if gc_start is not None:
                try:
                    lunch_start = int(gc_start.value)
                    lunch_row = (lunch_start - 480) // 60 + 1
                except (ValueError, TypeError):
                    pass

    # Set up lunch text (but delay the merge until after vertical merges to avoid python-docx bug)
    lunch_cell = table.cell(lunch_row, 1)
    lunch_cell.text = 'LUNCH BREAK'
    _set_cell_shading(lunch_cell, 'D9D9D9')
    for p in lunch_cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9)
            r.bold = True

    # Map sessions
    day_map = {'MONDAY': 1, 'TUESDAY': 2, 'WEDNESDAY': 3, 'THURSDAY': 4, 'FRIDAY': 5}
    
    unique_modules = {} # module_code -> {name, lecturers_set}
    unique_resources = {} # resource_name -> type
    
    for s in sessions:
        if s.module:
            if s.module.code not in unique_modules:
                unique_modules[s.module.code] = {
                    "name": s.module.name,
                    "lecturers": set()
                }
            if s.lecturer and s.lecturer.user:
                # Format: Prof. Name or Mr. Name based on something? Just use name
                unique_modules[s.module.code]["lecturers"].add(f"{s.lecturer.user.first_name} {s.lecturer.user.last_name}")
        
        if s.resource:
            unique_resources[s.resource.name] = s.resource.type

        d_idx = day_map.get(s.day_of_week)
        if not d_idx:
            continue
            
        start_mins = s.start_time.hour * 60 + s.start_time.minute
        end_mins = s.end_time.hour * 60 + s.end_time.minute
        duration_hours = max(1, (end_mins - start_mins) // 60)
        
        start_row = (start_mins - 480) // 60 + 1
        
        if start_row < 1 or start_row > 12:
            continue
            
        # Write to cell
        cell = table.cell(start_row, d_idx)
        mod_code = s.module.code if s.module else "Unknown"
        room_name = s.resource.name if s.resource else "Unknown"
        cell.text = f"{mod_code} ({room_name})"
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9)
                
        # Merge if duration > 1
        if duration_hours > 1:
            merge_base = cell
            for extra in range(1, duration_hours):
                next_row = start_row + extra
                if next_row > 12:
                    break
                if next_row == lunch_row:
                    merge_base = None
                    continue
                
                target_cell = table.cell(next_row, d_idx)
                if merge_base is None:
                    merge_base = target_cell
                    merge_base.text = f"{mod_code} ({room_name})"
                    for p in merge_base.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:
                            r.font.name = 'Times New Roman'
                            r.font.size = Pt(9)
                else:
                    try:
                        merge_base.merge(target_cell)
                    except Exception as e:
                        print(f"Failed to merge cells: {e}")

    # Now merge lunch row horizontally (doing this earlier causes ValueError in python-docx)
    try:
        table.cell(lunch_row, 1).merge(table.cell(lunch_row, 5))
    except Exception:
        pass

    doc.add_paragraph() # spacing

    # Table 2: Resources Legend
    if unique_resources:
        res_list = list(unique_resources.items())
        # Let's put them in 2 columns
        res_table = doc.add_table(rows=(len(res_list)+1)//2, cols=4)
        for idx, (r_name, r_type) in enumerate(res_list):
            row = idx // 2
            col_offset = (idx % 2) * 2
            cell_name = res_table.cell(row, col_offset)
            cell_type = res_table.cell(row, col_offset + 1)
            cell_name.text = r_name
            cell_type.text = r_type
            
            for p in cell_name.paragraphs + cell_type.paragraphs:
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(9)
        doc.add_paragraph()
    
    doc.add_paragraph() # spacing

    # Table 3: Courses List
    if unique_modules:
        course_table = doc.add_table(rows=len(unique_modules) + 1, cols=3)
        course_table.style = 'Table Grid'
        
        headers = ["Course Code", "Course Name", "Lecturer Panel"]
        for idx, h in enumerate(headers):
            cell = course_table.cell(0, idx)
            cell.text = h
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10)
                    r.bold = True
                    
        for r_idx, (code, info) in enumerate(sorted(unique_modules.items()), start=1):
            course_table.cell(r_idx, 0).text = code
            course_table.cell(r_idx, 1).text = info["name"]
            course_table.cell(r_idx, 2).text = " / ".join(info["lecturers"]) if info["lecturers"] else "TBD"
            
            for c_idx in range(3):
                cell = course_table.cell(r_idx, c_idx)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(9)

    doc.add_paragraph()
    doc.add_paragraph()
    
    # Signatures
    sig_table = doc.add_table(rows=3, cols=2)
    sig_table.cell(0, 0).text = f"Head of the Department (Acting)\nDepartment of {dept_name}\nDate"
    sig_table.cell(0, 1).text = f"Dean\nFaculty of {faculty_name}\nDate"
    
    for c_idx in range(2):
        cell = sig_table.cell(0, c_idx)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)

@router.get("/")
def export_timetable(
    batch_id: Optional[int] = None,
    degree_id: Optional[int] = None,
    dept_id: Optional[int] = None,
    faculty_id: Optional[int] = None,
    status: str = "PUBLISHED",
    format: str = "pdf",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    format_name = format.lower().strip()
    if format_name not in {"pdf", "docx"}:
        raise HTTPException(status_code=422, detail="format must be pdf or docx")
        
    _ensure_dependencies(format_name)
    
    from app.models.academic import Degree, Department
    query = db.query(Batch)
    
    if batch_id:
        query = query.filter(Batch.batch_id == batch_id)
    elif degree_id:
        query = query.filter(Batch.degree_id == degree_id)
    elif dept_id:
        query = query.join(Degree).filter(Degree.dept_id == dept_id)
    elif faculty_id:
        query = query.join(Degree).join(Department).filter(Department.faculty_id == faculty_id)
    else:
        raise HTTPException(status_code=400, detail="Provide at least one filter")
        
    batches = query.all()
    if not batches:
        raise HTTPException(status_code=404, detail="No batches found")
        
    import zipfile
    
    if len(batches) == 1:
        batch = batches[0]
        active_term = db.query(BatchActiveTerm).filter(BatchActiveTerm.batch_id == batch.batch_id, BatchActiveTerm.is_active == True).order_by(BatchActiveTerm.id.desc()).first()
        sessions = db.query(TimetableSession).filter(TimetableSession.batch_id == batch.batch_id, TimetableSession.status == status).all()
        if not sessions:
            raise HTTPException(status_code=404, detail="No timetable sessions found for the given criteria")
            
        doc = docx.Document()
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)
            
        append_timetable_to_doc(doc, batch, active_term, sessions, is_first=True)
        
        output = BytesIO()
        doc.save(output)
        docx_bytes = output.getvalue()
        
        semester_name = active_term.semester_name if active_term else f"Semester {batch.current_semester}"
        try:
            year = ((batch.current_semester - 1) // 2) + 1
            filename = f"{batch.batch_code} Year {year} {semester_name}.{format_name}"
        except Exception:
            filename = f"{batch.batch_code} {semester_name}.{format_name}"
            
        filename = "".join(c for c in filename if c.isalnum() or c in " ._-")
        
        if format_name == "docx":
            return StreamingResponse(
                BytesIO(docx_bytes),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f'attachment; filename="{filename}"'}
            )
        else:
            with tempfile.TemporaryDirectory() as temp_dir:
                docx_path = Path(temp_dir) / "timetable.docx"
                pdf_path = Path(temp_dir) / "timetable.pdf"
                docx_path.write_bytes(docx_bytes)
                
                try:
                    if sys.platform == "win32":
                        import pythoncom
                        pythoncom.CoInitialize()
                        if docx_to_pdf is None:
                            raise Exception("docx2pdf is not installed")
                        docx_to_pdf(str(docx_path), str(pdf_path))
                    else:
                        subprocess.run([
                            "libreoffice", "--headless", "--convert-to", "pdf",
                            "--outdir", str(temp_dir), str(docx_path)
                        ], check=True)
                except Exception as e:
                    raise HTTPException(status_code=500, detail=f"PDF conversion failed: {str(e)}")
                finally:
                    if sys.platform == "win32":
                        try:
                            import pythoncom
                            pythoncom.CoUninitialize()
                        except Exception:
                            pass

                if not pdf_path.exists():
                    raise HTTPException(status_code=500, detail="Failed to generate PDF")
                return StreamingResponse(
                    BytesIO(pdf_path.read_bytes()),
                    media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'}
                )
    
    # Bulk export -> ZIP file
    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        added_count = 0
        with tempfile.TemporaryDirectory() as temp_dir:
            for batch in batches:
                active_term = db.query(BatchActiveTerm).filter(BatchActiveTerm.batch_id == batch.batch_id, BatchActiveTerm.is_active == True).order_by(BatchActiveTerm.id.desc()).first()
                sessions = db.query(TimetableSession).filter(TimetableSession.batch_id == batch.batch_id, TimetableSession.status == status).all()
                if not sessions:
                    continue
                    
                doc = docx.Document()
                for section in doc.sections:
                    section.top_margin = Inches(0.5)
                    section.bottom_margin = Inches(0.5)
                    section.left_margin = Inches(0.5)
                    section.right_margin = Inches(0.5)
                    
                append_timetable_to_doc(doc, batch, active_term, sessions, is_first=True)
                
                output = BytesIO()
                doc.save(output)
                docx_bytes = output.getvalue()
                
                semester_name = active_term.semester_name if active_term else f"Semester {batch.current_semester}"
                try:
                    year = ((batch.current_semester - 1) // 2) + 1
                    batch_filename = f"{batch.batch_code} Year {year} {semester_name}.{format_name}"
                except Exception:
                    batch_filename = f"{batch.batch_code} {semester_name}.{format_name}"
                batch_filename = "".join(c for c in batch_filename if c.isalnum() or c in " ._-")
                
                if format_name == "docx":
                    zip_file.writestr(batch_filename, docx_bytes)
                    added_count += 1
                else:
                    docx_path = Path(temp_dir) / f"{batch.batch_code}.docx"
                    pdf_path = Path(temp_dir) / f"{batch.batch_code}.pdf"
                    docx_path.write_bytes(docx_bytes)
                    
                    try:
                        if sys.platform == "win32":
                            import pythoncom
                            pythoncom.CoInitialize()
                            if docx_to_pdf is not None:
                                docx_to_pdf(str(docx_path), str(pdf_path))
                        else:
                            subprocess.run([
                                "libreoffice", "--headless", "--convert-to", "pdf",
                                "--outdir", str(temp_dir), str(docx_path)
                            ], check=True)
                            
                        if pdf_path.exists():
                            zip_file.write(pdf_path, batch_filename)
                            added_count += 1
                    except Exception as e:
                        print(f"Failed to convert {batch.batch_code}: {e}")
                    finally:
                        if sys.platform == "win32":
                            try:
                                import pythoncom
                                pythoncom.CoUninitialize()
                            except Exception:
                                pass

        if added_count == 0:
            raise HTTPException(status_code=404, detail="No timetable sessions found to export")
            
    zip_buffer.seek(0)
    degree_code = batches[0].degree.code if batches[0].degree else "Timetables"
    zip_filename = f"{degree_code}_Timetables.zip"
    zip_filename = "".join(c for c in zip_filename if c.isalnum() or c in " ._-")
    
    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{zip_filename}"'}
    )
