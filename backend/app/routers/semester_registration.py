from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Optional
import tempfile

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.database.connection import get_db
from app.models.academic import Batch, BatchActiveTerm, DegreeSemesterModule, Module
from app.models.user import User, UserRole
from app.utils.dependencies import require_roles

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.table import WD_ROW_HEIGHT_RULE
except ImportError:  # pragma: no cover - dependency injected via requirements
    Document = None
    Pt = None
    Inches = None
    WD_ROW_HEIGHT_RULE = None

try:
    from docx2pdf import convert as docx_to_pdf
except ImportError:  # pragma: no cover - dependency injected via requirements
    docx_to_pdf = None


router = APIRouter(prefix="/api/semester-registration", tags=["semester-registration"])


class SemesterRegistrationPayload(BaseModel):
    enrollment_no: Optional[str] = None
    faculty_name: Optional[str] = None
    course_of_study: Optional[str] = None
    academic_year: Optional[str] = None
    full_name: str = Field(min_length=1)
    name_with_initials: str = Field(min_length=1)
    postal_address: str = Field(min_length=1)
    phone: Optional[str] = None
    mobile: Optional[str] = None
    email: Optional[str] = None
    scholarship_name: Optional[str] = None
    scholarship_amount: Optional[str] = None
    selected_module_ids: list[int] = Field(default_factory=list)
    form_date: Optional[str] = None


TEMPLATE_PATH = (
    Path(__file__).resolve().parents[3]
    / "frontend"
    / "public"
    / "semester-registration.docx"
)
SAMPLE_ENROLLMENT_NO = "UWU/ICT/21/010"
SAMPLE_FACULTY = "Faculty of Technological Studies"
SAMPLE_COURSE = "Bachelor of Information and Communication Technology Honours (BICT)"
SAMPLE_FULL_NAME = "Wickramalage Nadeesh Malaka Chathuranga"
SAMPLE_INITIALS = "W.N.M Chathuranga"
SAMPLE_ADDRESS = "No:46/3,Kodurawa,Polgasowita."
SAMPLE_PHONE = "0787620370"
SAMPLE_MOBILE = "0774902773"
SAMPLE_EMAIL = "nadeeshmalaka2001@gmail.com"
SAMPLE_DATE = "8/30/2022"
SAMPLE_LEVEL = "100 Level"
SAMPLE_LEVEL_ALT = "400 Level"
SAMPLE_SEMESTER = "1st Year 1st Semester"
SAMPLE_SEMESTER_IN_YEAR = "1st Semester"
SAMPLE_ACADEMIC_YEAR = "2022/2023"
SAMPLE_BATCH_YEAR = "Batch Year 2020/2021"
SAMPLE_SEMESTER_FULL = "1st Year 1st Semester, Academic Year 2022/2023"
SAMPLE_ACADEMIC_YEAR_FULL = "Academic Year 2022/2023"


def _semester_label(semester_number: int) -> str:
    year = ((semester_number - 1) // 2) + 1
    semester = 1 if semester_number % 2 == 1 else 2
    return f"Year {year} Semester {semester}"


SEMESTER_NAME_TO_NUMBER = {
    _semester_label(number): number
    for number in range(1, 11)
}


def _resolve_batch_semester(batch: Batch, db: Session) -> tuple[int, str, Optional[str]]:
    active_term = (
        db.query(BatchActiveTerm)
        .filter(
            BatchActiveTerm.batch_id == batch.batch_id,
            BatchActiveTerm.is_active.is_(True),
        )
        .order_by(BatchActiveTerm.id.desc())
        .first()
    )

    academic_year = None
    if active_term:
        academic_year = getattr(active_term, "academic_year", None)

    if active_term and active_term.semester_name in SEMESTER_NAME_TO_NUMBER:
        semester_number = SEMESTER_NAME_TO_NUMBER[active_term.semester_name]
        semester_name = active_term.semester_name
    else:
        semester_number = batch.current_semester
        semester_name = _semester_label(semester_number)

    return semester_number, semester_name, academic_year


def _get_student_batch(current_user: User, db: Session) -> Batch:
    student_profile = current_user.student_profile
    if not student_profile:
        raise HTTPException(status_code=404, detail="Student profile not found")

    raw_batch_id = getattr(student_profile, "batch_id", None)
    if raw_batch_id is None:
        raw_batch_id = getattr(student_profile, "batch", None)

    try:
        batch_id = int(str(raw_batch_id).strip()) if raw_batch_id is not None else None
    except (TypeError, ValueError):
        batch_id = None

    if not batch_id:
        raise HTTPException(status_code=422, detail="Student batch is not set")

    batch = db.query(Batch).filter(Batch.batch_id == batch_id).first()
    if not batch:
        raise HTTPException(status_code=404, detail="Batch not found")

    return batch


def _get_current_semester_modules(batch: Batch, semester_number: int, db: Session) -> list[Module]:
    return (
        db.query(Module)
        .join(DegreeSemesterModule, DegreeSemesterModule.module_id == Module.module_id)
        .filter(
            DegreeSemesterModule.degree_id == batch.degree_id,
            DegreeSemesterModule.semester_number == semester_number,
        )
        .order_by(Module.code.asc())
        .all()
    )


def _format_level_label(semester_number: int) -> str:
    year = ((semester_number - 1) // 2) + 1
    return f"{year * 100} Level"


def _ordinal(value: int) -> str:
    if 10 <= value % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(value % 10, "th")
    return f"{value}{suffix}"


def _format_semester_title(semester_number: int) -> str:
    year = ((semester_number - 1) // 2) + 1
    semester_in_year = 1 if semester_number % 2 == 1 else 2
    return f"{_ordinal(year)} Year {_ordinal(semester_in_year)} Semester"


@router.get("/form-data")
def get_form_data(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.STUDENT)),
):
    batch = _get_student_batch(current_user, db)
    if not batch.degree:
        raise HTTPException(status_code=404, detail="Degree not found for student batch")

    semester_number, semester_name, academic_year = _resolve_batch_semester(batch, db)
    modules = _get_current_semester_modules(batch, semester_number, db)

    faculty_name = None
    department_name = None
    if batch.degree.department:
        department_name = batch.degree.department.name
        if batch.degree.department.faculty:
            faculty_name = batch.degree.department.faculty.name

    reg_no = getattr(current_user.student_profile, "reg_no", None) or getattr(
        current_user.student_profile, "index_number", None
    )
    registration_number = getattr(current_user.student_profile, "registration_number", None)

    return {
        "student": {
            "full_name": f"{current_user.first_name} {current_user.last_name}".strip(),
            "reg_no": reg_no,
            "registration_number": registration_number,
            "email": current_user.email,
            "contact_number": current_user.contact_number,
        },
        "batch": {
            "batch_id": batch.batch_id,
            "batch_code": batch.batch_code,
        },
        "degree": {
            "degree_id": batch.degree_id,
            "degree_name": batch.degree.name,
        },
        "faculty": {
            "name": faculty_name,
        },
        "department": {
            "name": department_name,
        },
        "semester": {
            "semester_number": semester_number,
            "semester_name": semester_name,
            "academic_year": academic_year,
            "level_label": _format_level_label(semester_number),
        },
        "modules": [
            {
                "module_id": module.module_id,
                "code": module.code,
                "name": module.name,
            }
            for module in modules
        ],
    }


def _ensure_dependencies(format_name: str) -> None:
    if format_name == "docx" and Document is None:
        raise HTTPException(status_code=500, detail="python-docx is not installed")
    if format_name == "pdf" and docx_to_pdf is None:
        raise HTTPException(status_code=500, detail="docx2pdf is not installed")


def _normalize(text: str) -> str:
    return " ".join(text.split()).strip().lower()


def _iter_document_paragraphs(document):
    yield from document.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _iter_document_tables(document):
    yield from document.tables
    for section in document.sections:
        yield from section.header.tables
        yield from section.footer.tables


def _replace_text_in_paragraph(
    paragraph,
    replacements: dict[str, str],
    underline_keys: set[str],
) -> None:
    if not paragraph.runs:
        return

    paragraph_text = "".join(run.text for run in paragraph.runs)
    updated_text = paragraph_text
    underline = False
    for old, new in replacements.items():
        if old in updated_text:
            updated_text = updated_text.replace(old, new)
            if old in underline_keys:
                underline = True

    if updated_text == paragraph_text:
        return

    paragraph.text = updated_text
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        if underline:
            run.underline = True


def _replace_in_document(
    document: Document,
    replacements: dict[str, str],
    underline_keys: set[str],
) -> None:
    for paragraph in _iter_document_paragraphs(document):
        _replace_text_in_paragraph(paragraph, replacements, underline_keys)

    for table in _iter_document_tables(document):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_text_in_paragraph(paragraph, replacements, underline_keys)


def _set_document_font(document):
    for paragraph in _iter_document_paragraphs(document):
        for run in paragraph.runs:
            run.font.name = "Times New Roman"

    for table in _iter_document_tables(document):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"


def _update_table_cell(
    table,
    label_keywords: list[str],
    value: str,
    underline: bool = False,
) -> bool:
    for row in table.rows:
        label_text = _normalize(row.cells[0].text) if row.cells else ""
        if all(keyword in label_text for keyword in label_keywords):
            if len(row.cells) > 1:
                row.cells[1].text = value
                if underline:
                    for paragraph in row.cells[1].paragraphs:
                        for run in paragraph.runs:
                            run.underline = True
            return True
    return False


def _update_module_table(document: Document, modules: list[Module]) -> None:
    for table in document.tables:
        if not table.rows:
            continue
        header_cells = [_normalize(cell.text) for cell in table.rows[0].cells]
        if "subject" in " ".join(header_cells) and "subject code" in " ".join(header_cells):
            for index, row in enumerate(table.rows[1:], start=1):
                if WD_ROW_HEIGHT_RULE is not None and Inches is not None:
                    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                    row.height = Inches(0.3)
                row.cells[0].text = f"{index:02d}."
                if index <= len(modules):
                    row.cells[1].text = modules[index - 1].name
                    row.cells[2].text = modules[index - 1].code
                else:
                    row.cells[1].text = ""
                    row.cells[2].text = ""

                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            if Pt is not None:
                                run.font.size = Pt(9)
                            run.font.name = "Times New Roman"
            break


def _apply_run_style(source_run, target_run) -> None:
    if not source_run:
        return
    target_run.bold = source_run.bold
    target_run.italic = source_run.italic
    target_run.font.name = "Times New Roman"
    target_run.font.size = source_run.font.size


def _reset_paragraph(paragraph) -> Optional[object]:
    base_run = paragraph.runs[0] if paragraph.runs else None
    paragraph.text = ""
    return base_run


def _set_paragraph_with_value(paragraph, label: str, value: str) -> None:
    base_run = _reset_paragraph(paragraph)
    label_run = paragraph.add_run(label)
    _apply_run_style(base_run, label_run)
    value_run = paragraph.add_run(f" {value}")
    _apply_run_style(base_run, value_run)
    value_run.underline = True


def _set_contact_paragraph(paragraph, phone: str, mobile: str, email: str) -> None:
    base_run = _reset_paragraph(paragraph)
    label_run = paragraph.add_run("Contact: Home:")
    _apply_run_style(base_run, label_run)
    home_run = paragraph.add_run(f" {phone}")
    _apply_run_style(base_run, home_run)
    home_run.underline = True
    mobile_label = paragraph.add_run("   Mobile:")
    _apply_run_style(base_run, mobile_label)
    mobile_run = paragraph.add_run(f" {mobile}")
    _apply_run_style(base_run, mobile_run)
    mobile_run.underline = True
    email_label = paragraph.add_run("   e-mail:")
    _apply_run_style(base_run, email_label)
    email_run = paragraph.add_run(f" {email}")
    _apply_run_style(base_run, email_run)
    email_run.underline = True


def _set_mobile_email_paragraph(paragraph, mobile: str, email: str) -> None:
    base_run = _reset_paragraph(paragraph)
    label_run = paragraph.add_run("Mobile:")
    _apply_run_style(base_run, label_run)
    mobile_run = paragraph.add_run(f" {mobile}")
    _apply_run_style(base_run, mobile_run)
    mobile_run.underline = True
    email_label = paragraph.add_run("   e-mail:")
    _apply_run_style(base_run, email_label)
    email_run = paragraph.add_run(f" {email}")
    _apply_run_style(base_run, email_run)
    email_run.underline = True


def _fill_paragraphs(document: Document, data: dict) -> None:
    date_filled = False
    for paragraph in document.paragraphs:
        normalized = _normalize(paragraph.text)

        if normalized.startswith("full name"):
            _set_paragraph_with_value(paragraph, "Full Name (Mr./ Miss.):", data["full_name"])
            continue

        if normalized.startswith("name with initials"):
            _set_paragraph_with_value(paragraph, "Name with Initials:", data["name_with_initials"])
            continue

        if normalized.startswith("postal address"):
            _set_paragraph_with_value(paragraph, "Postal Address:", data["postal_address"])
            continue

        if normalized.startswith("contact: home") and "mobile" in normalized and "e-mail" in normalized:
            _set_contact_paragraph(paragraph, data["phone"], data["mobile"], data["email"])
            continue

        if normalized.startswith("contact: home") and "mobile" not in normalized:
            base_run = _reset_paragraph(paragraph)
            label_run = paragraph.add_run("Contact: Home:")
            _apply_run_style(base_run, label_run)
            home_run = paragraph.add_run(f" {data['phone']}")
            _apply_run_style(base_run, home_run)
            home_run.underline = True
            continue

        if normalized.startswith("mobile") and "e-mail" in normalized:
            _set_mobile_email_paragraph(paragraph, data["mobile"], data["email"])
            continue

        if normalized.startswith("date:") and "signature" not in normalized:
            if not date_filled:
                _set_paragraph_with_value(paragraph, "Date:", data["form_date"])
                date_filled = True
            continue

        if normalized.startswith("enrollment no"):
            _set_paragraph_with_value(paragraph, "Enrollment No:", data["registration_no"])
            continue

        if normalized.startswith("scholarship"):
            _set_paragraph_with_value(
                paragraph,
                "Scholarship: (i.) Name/Source:",
                data["scholarship_name"],
            )
            continue

        if normalized.startswith("(ii.) annual payment"):
            _set_paragraph_with_value(
                paragraph,
                "(ii.) Annual Payment:",
                data["scholarship_amount"],
            )
            continue


def _render_docx(data: dict, modules: list[Module]) -> bytes:
    if Document is None:
        raise HTTPException(status_code=500, detail="python-docx is not installed")
    if not TEMPLATE_PATH.exists():
        raise HTTPException(status_code=500, detail="DOCX template not found")

    document = Document(TEMPLATE_PATH)

    replacements = {
        SAMPLE_ENROLLMENT_NO: data["registration_no"],
        SAMPLE_FACULTY: data["faculty"],
        SAMPLE_COURSE: data["course_of_study"],
        SAMPLE_FULL_NAME: data["full_name"],
        SAMPLE_INITIALS: data["name_with_initials"],
        SAMPLE_ADDRESS: data["postal_address"],
        SAMPLE_PHONE: data["phone"],
        SAMPLE_MOBILE: data["mobile"],
        SAMPLE_EMAIL: data["email"],
        SAMPLE_DATE: data["form_date"],
        SAMPLE_LEVEL: data["level_label"],
        SAMPLE_LEVEL_ALT: data["level_label"],
        SAMPLE_SEMESTER: data["semester_label"],
        SAMPLE_SEMESTER_IN_YEAR: data["semester_in_year_label"],
        SAMPLE_SEMESTER_FULL: f"{data['semester_label']}, Academic Year {data['academic_year']}",
        SAMPLE_ACADEMIC_YEAR: data["academic_year"],
        SAMPLE_ACADEMIC_YEAR_FULL: f"Academic Year {data['academic_year']}",
        SAMPLE_BATCH_YEAR: data["batch_year_label"],
    }

    underline_keys = {
        SAMPLE_ENROLLMENT_NO,
        SAMPLE_FACULTY,
        SAMPLE_COURSE,
        SAMPLE_FULL_NAME,
        SAMPLE_INITIALS,
        SAMPLE_ADDRESS,
        SAMPLE_PHONE,
        SAMPLE_MOBILE,
        SAMPLE_EMAIL,
        SAMPLE_DATE,
        SAMPLE_ACADEMIC_YEAR,
        SAMPLE_ACADEMIC_YEAR_FULL,
        SAMPLE_BATCH_YEAR,
    }

    _replace_in_document(document, replacements, underline_keys)

    _fill_paragraphs(document, data)

    for table in document.tables:
        _update_table_cell(table, ["enrollment"], data["registration_no"], underline=True)
        _update_table_cell(table, ["faculty"], data["faculty"], underline=True)
        _update_table_cell(table, ["course of study"], data["course_of_study"], underline=True)
        _update_table_cell(table, ["name with initials"], data["name_with_initials"], underline=True)
        _update_table_cell(table, ["postal address"], data["postal_address"], underline=True)
        _update_table_cell(table, ["scholarship", "name"], data["scholarship_name"], underline=True)
        _update_table_cell(table, ["annual payment"], data["scholarship_amount"], underline=True)
        _update_table_cell(table, ["date"], data["form_date"], underline=True)

        if _update_table_cell(
            table,
            ["contact", "phone"],
            f"{data['phone']} / {data['mobile']} / {data['email']}",
            underline=True,
        ):
            continue

    _update_module_table(document, modules)
    _set_document_font(document)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _render_pdf(data: dict, modules: list[Module]) -> bytes:
    if docx_to_pdf is None:
        raise HTTPException(status_code=500, detail="docx2pdf is not installed")

    docx_bytes = _render_docx(data, modules)

    with tempfile.TemporaryDirectory() as temp_dir:
        docx_path = Path(temp_dir) / "semester_registration.docx"
        pdf_path = Path(temp_dir) / "semester_registration.pdf"
        docx_path.write_bytes(docx_bytes)
        
        try:
            import pythoncom
            pythoncom.CoInitialize()
            docx_to_pdf(str(docx_path), str(pdf_path))
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"PDF conversion failed: {str(e)}")
        finally:
            try:
                import pythoncom
                pythoncom.CoUninitialize()
            except Exception:
                pass

        if not pdf_path.exists():
            raise HTTPException(status_code=500, detail="Failed to generate PDF")
        return pdf_path.read_bytes()


@router.post("/generate")
def generate_form(
    payload: SemesterRegistrationPayload,
    format: str = "pdf",
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.STUDENT)),
):
    format_name = format.lower().strip()
    if format_name not in {"pdf", "docx"}:
        raise HTTPException(status_code=422, detail="format must be pdf or docx")

    _ensure_dependencies(format_name)

    batch = _get_student_batch(current_user, db)
    if not batch.degree:
        raise HTTPException(status_code=404, detail="Degree not found for student batch")

    semester_number, semester_name, academic_year = _resolve_batch_semester(batch, db)
    modules = _get_current_semester_modules(batch, semester_number, db)

    module_lookup = {module.module_id: module for module in modules}
    selected_modules = [
        module_lookup[module_id]
        for module_id in payload.selected_module_ids
        if module_id in module_lookup
    ]
    if not selected_modules:
        raise HTTPException(status_code=422, detail="Select at least one module")

    faculty_name = "-"
    if batch.degree.department and batch.degree.department.faculty:
        faculty_name = batch.degree.department.faculty.name

    registration_no = (
        payload.enrollment_no
        or getattr(current_user.student_profile, "registration_number", None)
        or getattr(current_user.student_profile, "reg_no", None)
        or getattr(current_user.student_profile, "index_number", None)
        or "-"
    )

    if payload.form_date:
        try:
            parsed_date = datetime.strptime(payload.form_date, "%Y-%m-%d")
            form_date = f"{parsed_date.month}/{parsed_date.day}/{parsed_date.year}"
        except ValueError:
            form_date = payload.form_date
    else:
        now = datetime.now()
        form_date = f"{now.month}/{now.day}/{now.year}"
    resolved_academic_year = payload.academic_year or academic_year or "-"
    resolved_faculty = payload.faculty_name or faculty_name or "-"
    resolved_course = payload.course_of_study or batch.degree.name
    semester_title = _format_semester_title(semester_number)
    semester_in_year = "1st Semester" if semester_number % 2 == 1 else "2nd Semester"
    batch_year_label = f"Batch Year {resolved_academic_year}"

    data = {
        "registration_no": registration_no,
        "faculty": resolved_faculty,
        "course_of_study": resolved_course,
        "full_name": payload.full_name.strip(),
        "name_with_initials": payload.name_with_initials.strip(),
        "postal_address": payload.postal_address.strip(),
        "phone": payload.phone or "-",
        "mobile": payload.mobile or "-",
        "email": payload.email or current_user.email,
        "semester_label": semester_title,
        "semester_in_year_label": semester_in_year,
        "academic_year": resolved_academic_year,
        "level_label": _format_level_label(semester_number),
        "batch_year_label": batch_year_label,
        "batch_code": batch.batch_code,
        "scholarship_name": payload.scholarship_name or "N/A",
        "scholarship_amount": payload.scholarship_amount or "N/A",
        "form_date": form_date,
    }

    if format_name == "docx":
        file_bytes = _render_docx(data, selected_modules)
        filename = "semester-registration.docx"
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        file_bytes = _render_pdf(data, selected_modules)
        filename = "semester-registration.pdf"
        media_type = "application/pdf"

    return StreamingResponse(
        BytesIO(file_bytes),
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

